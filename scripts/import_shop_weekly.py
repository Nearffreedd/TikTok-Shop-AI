"""
店铺周数据导入脚本
功能：读取店铺周数据 Excel -> 写入 SQLite 数据库
用法：python import_shop_weekly.py <Excel文件路径> [店铺代码]

店铺数据表结构（第3行为表头，第4行为数据）：
- 第1行：分析日期范围
- 第2行：数据概览标题
- 第3行：表头（28列）
- 第4行：总计值（本期数据）
- 第5行：百分比变化（环比）
"""

import sqlite3
import openpyxl
import os
import sys
import re
from datetime import datetime

DB_PATH = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                       'Layer1_Permanent', '01_My_Shop', 'shop_data.db')


def parse_week_range_from_row(row1_text):
    """从第1行解析日期范围，如 '分析日期：18/05/2026-24/05/2026' -> ('2026-05-18', '2026-05-24')"""
    match = re.search(r'(\d{2})/(\d{2})/(\d{4})-(\d{2})/(\d{2})/(\d{4})', str(row1_text))
    if match:
        start = f"{match.group(3)}-{match.group(2)}-{match.group(1)}"
        end = f"{match.group(6)}-{match.group(5)}-{match.group(4)}"
        return start, end
    return None, None


def detect_shop_code(filename):
    match = re.match(r'([A-Z]+)-店铺数据', filename)
    if match:
        return match.group(1)
    return None


def parse_float(val):
    """安全解析数值，处理 ₱ 符号和千位分隔符"""
    if val is None or val == '' or val == '-':
        return 0.0
    try:
        s = str(val).strip()
        # 移除货币符号（₱、PHP、$ 等）、千位分隔符和百分号
        s = s.replace('PHP', '').replace('₱', '').replace('$', '').replace(',', '').replace('%', '').strip()
        return float(s)
    except:
        return 0.0


def import_shop_data(filepath, shop_code=None):
    if not os.path.exists(filepath):
        print(f'[X] 文件不存在: {filepath}')
        return False

    filename = os.path.basename(filepath)

    if shop_code is None:
        shop_code = detect_shop_code(filename)
        if shop_code is None:
            print('[X] 无法从文件名识别店铺代码')
            return False

    wb = openpyxl.load_workbook(filepath)
    ws = wb.active

    # 读取第1行获取日期
    row1 = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    week_start, week_end = parse_week_range_from_row(row1[0] if row1 else '')
    if week_start is None:
        print('[X] 无法从第1行解析日期范围')
        return False

    # 读取表头（第3行）和数据（第4行）
    headers = [c.value for c in next(ws.iter_rows(min_row=3, max_row=3))]
    data_row = [c.value for c in next(ws.iter_rows(min_row=4, max_row=4))]
    change_row = [c.value for c in next(ws.iter_rows(min_row=5, max_row=5))]

    d = dict(zip(headers, data_row))
    c = dict(zip(headers, change_row))

    print(f'[文件] {filename}')
    print(f'[店铺] {shop_code}')
    print(f'[周期] {week_start} ~ {week_end}')

    conn = sqlite3.connect(DB_PATH)
    cursor = conn.cursor()

    # 获取店铺ID
    cursor.execute('SELECT shop_id FROM shops WHERE shop_code = ?', (shop_code,))
    row = cursor.fetchone()
    if row is None:
        print(f'[X] 店铺 {shop_code} 不存在，请先运行 init_database.py')
        conn.close()
        return False
    shop_id = row[0]

    # 检查表是否存在，不存在则创建
    cursor.execute('''
        CREATE TABLE IF NOT EXISTS shop_weekly (
            record_id INTEGER PRIMARY KEY AUTOINCREMENT,
            shop_id INTEGER NOT NULL,
            week_start TEXT NOT NULL,
            week_end TEXT NOT NULL,

            -- 核心指标
            gmv REAL DEFAULT 0,
            orders INTEGER DEFAULT 0,
            customers INTEGER DEFAULT 0,
            units_sold INTEGER DEFAULT 0,
            refund_units REAL DEFAULT 0,
            sku_orders INTEGER DEFAULT 0,
            total_revenue REAL DEFAULT 0,
            page_views INTEGER DEFAULT 0,
            visitors INTEGER DEFAULT 0,
            conversion_rate REAL DEFAULT 0,
            product_impressions INTEGER DEFAULT 0,
            deduped_product_impressions INTEGER DEFAULT 0,
            product_clicks INTEGER DEFAULT 0,
            deduped_clicks INTEGER DEFAULT 0,
            avg_order_value REAL DEFAULT 0,

            -- 达人直播
            affiliate_live_attributed_gmv REAL DEFAULT 0,
            affiliate_live_gmv REAL DEFAULT 0,
            affiliate_live_indirect_gmv REAL DEFAULT 0,
            bound_account_live_attributed_gmv REAL DEFAULT 0,
            merchant_live_gmv REAL DEFAULT 0,
            merchant_live_indirect_gmv REAL DEFAULT 0,

            -- 联盟视频
            affiliate_video_attributed_gmv REAL DEFAULT 0,
            affiliate_video_gmv REAL DEFAULT 0,
            affiliate_video_indirect_gmv REAL DEFAULT 0,
            bound_account_video_attributed_gmv REAL DEFAULT 0,
            merchant_video_gmv REAL DEFAULT 0,
            merchant_video_indirect_gmv REAL DEFAULT 0,

            -- 环比变化
            gmv_change REAL DEFAULT 0,
            orders_change REAL DEFAULT 0,
            customers_change REAL DEFAULT 0,
            units_sold_change REAL DEFAULT 0,
            conversion_rate_change REAL DEFAULT 0,

            source_file TEXT,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,

            FOREIGN KEY (shop_id) REFERENCES shops(shop_id),
            UNIQUE(shop_id, week_start, week_end)
        )
    ''')

    # 解析数值
    gmv = parse_float(d.get('GMV'))
    orders = int(parse_float(d.get('订单数')))
    customers = int(parse_float(d.get('客户数')))
    units = int(parse_float(d.get('商品成交件数')))
    refund_units = parse_float(d.get('已退款的商品件数'))
    sku_orders = int(parse_float(d.get('SKU 订单数')))
    total_revenue = parse_float(d.get('总成交额'))
    page_views = int(parse_float(d.get('页面浏览次数')))
    visitors = int(parse_float(d.get('商品访客数')))
    cvr = parse_float(d.get('转化率'))
    prod_imp = int(parse_float(d.get('商品曝光次数')))
    deduped_imp = int(parse_float(d.get('去重商品曝光次 数')))
    prod_clicks = int(parse_float(d.get('商品点击量')))
    deduped_clicks = int(parse_float(d.get('去重点击次数')))
    aov = parse_float(d.get('平均订单金额'))

    # 达人直播
    aff_live_attr = parse_float(d.get('达人直播归因 GMV'))
    aff_live = parse_float(d.get('达人直播 GMV'))
    aff_live_indirect = parse_float(d.get('达人直播间接 GMV'))
    bound_live_attr = parse_float(d.get('绑定账号直播归因 GMV'))
    merch_live = parse_float(d.get('商家直播 GMV'))
    merch_live_indirect = parse_float(d.get('商家直播间接 GMV'))

    # 联盟视频
    aff_video_attr = parse_float(d.get('联盟视频归因 GMV'))
    aff_video = parse_float(d.get('达人视频 GMV'))
    aff_video_indirect = parse_float(d.get('达人视频间接 GMV'))
    bound_video_attr = parse_float(d.get('绑定账号视频归因 GMV'))
    merch_video = parse_float(d.get('商家视频 GMV'))
    merch_video_indirect = parse_float(d.get('商家视频间接 GMV'))

    # 环比变化
    gmv_change = parse_float(c.get('GMV'))
    orders_change = parse_float(c.get('订单数'))
    customers_change = parse_float(c.get('客户数'))
    units_change = parse_float(c.get('商品成交件数'))
    cvr_change = parse_float(c.get('转化率'))

    try:
        cursor.execute('''
            INSERT INTO shop_weekly (
                shop_id, week_start, week_end,
                gmv, orders, customers, units_sold, refund_units, sku_orders,
                total_revenue, page_views, visitors, conversion_rate,
                product_impressions, deduped_product_impressions,
                product_clicks, deduped_clicks, avg_order_value,
                affiliate_live_attributed_gmv, affiliate_live_gmv, affiliate_live_indirect_gmv,
                bound_account_live_attributed_gmv, merchant_live_gmv, merchant_live_indirect_gmv,
                affiliate_video_attributed_gmv, affiliate_video_gmv, affiliate_video_indirect_gmv,
                bound_account_video_attributed_gmv, merchant_video_gmv, merchant_video_indirect_gmv,
                gmv_change, orders_change, customers_change, units_sold_change, conversion_rate_change,
                source_file
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        ''', (
            shop_id, week_start, week_end,
            gmv, orders, customers, units, refund_units, sku_orders,
            total_revenue, page_views, visitors, cvr,
            prod_imp, deduped_imp, prod_clicks, deduped_clicks, aov,
            aff_live_attr, aff_live, aff_live_indirect,
            bound_live_attr, merch_live, merch_live_indirect,
            aff_video_attr, aff_video, aff_video_indirect,
            bound_video_attr, merch_video, merch_video_indirect,
            gmv_change, orders_change, customers_change, units_change, cvr_change,
            filename
        ))
        print(f'[完成] 店铺数据导入完成')
    except sqlite3.IntegrityError:
        print(f'[跳过] 该周期数据已存在，跳过')

    conn.commit()
    conn.close()
    return True


if __name__ == '__main__':
    if len(sys.argv) < 2:
        print('用法: python import_shop_weekly.py <Excel文件路径> [店铺代码]')
        sys.exit(1)

    filepath = sys.argv[1]
    shop_code = sys.argv[2] if len(sys.argv) > 2 else None

    if not os.path.isabs(filepath):
        filepath = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), filepath)

    import_shop_data(filepath, shop_code)
