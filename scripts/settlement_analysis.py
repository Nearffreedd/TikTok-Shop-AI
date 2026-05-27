import openpyxl
import os
import sys
import io
from datetime import datetime, timedelta
from collections import defaultdict, Counter

# 设置输出编码为UTF-8
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

filepath = os.path.join(os.path.dirname(os.path.dirname(__file__)), 'Layer2_Working', '结算统计.xlsx')
wb = openpyxl.load_workbook(filepath, data_only=True)

# ========== 读取数据 ==========
ws = wb[wb.sheetnames[0]]
rows = list(ws.iter_rows(min_row=2, values_only=True))  # 跳过表头

# 解析列索引（根据观察到的数据）
# 0:订单ID, 1:状态, 2:退款金额, 3:商品总额, 4:下单时间, 5:结算时间, 
# 6:取消方, 7:结算金额, 8:商品总额(重复), 9:平台费用, 10:佣金, 
# 11:运费, 12:平台优惠券, 13:达人佣金, 14:其他扣费, 15:其他扣费2

completed = []
cancelled = []
for r in rows:
    status = r[1]
    if status == 'Completed':
        completed.append(r)
    elif status == 'Cancelled':
        cancelled.append(r)

print(f"总订单数: {len(rows)}")
print(f"已完成订单: {len(completed)}")
print(f"已取消订单: {len(cancelled)}")
print()

# ========== 1. 整体结算概况 ==========
print("=" * 70)
print("一、整体结算概况")
print("=" * 70)

# 已完成订单的结算金额
settlement_amounts = []
for r in completed:
    amt = r[7]  # 结算金额
    if amt is not None and isinstance(amt, (int, float)):
        settlement_amounts.append(amt)

total_settlement = sum(settlement_amounts)
avg_settlement = total_settlement / len(settlement_amounts) if settlement_amounts else 0
print(f"已完成订单数: {len(completed)}")
print(f"总结算金额: ₱{total_settlement:,.2f}")
print(f"平均每单结算金额: ₱{avg_settlement:,.2f}")
print(f"结算金额范围: ₱{min(settlement_amounts):,.2f} ~ ₱{max(settlement_amounts):,.2f}")
print()

# ========== 2. 每日结算分析 ==========
print("=" * 70)
print("二、每日结算分析（按结算日期）")
print("=" * 70)

daily_settlement = defaultdict(lambda: {'count': 0, 'amount': 0, 'gmv': 0})
for r in completed:
    settle_time = r[5]  # 结算时间
    if settle_time is None:
        continue
    if isinstance(settle_time, datetime):
        date_key = settle_time.strftime('%Y-%m-%d')
    else:
        continue
    amt = r[7] if r[7] is not None else 0
    gmv = r[3] if r[3] is not None else 0
    daily_settlement[date_key]['count'] += 1
    daily_settlement[date_key]['amount'] += amt
    daily_settlement[date_key]['gmv'] += gmv

print(f"{'日期':<14} {'订单数':<8} {'结算金额':<14} {'商品总额(GMV)':<16} {'结算率':<10}")
print("-" * 62)
for date_key in sorted(daily_settlement.keys()):
    d = daily_settlement[date_key]
    rate = d['amount'] / d['gmv'] * 100 if d['gmv'] > 0 else 0
    print(f"{date_key:<14} {d['count']:<8} ₱{d['amount']:<10,.2f} ₱{d['gmv']:<12,.2f} {rate:<9.1f}%")

print()

# ========== 3. 每日下单分析 ==========
print("=" * 70)
print("三、每日下单分析（按下单日期）")
print("=" * 70)

daily_orders = defaultdict(lambda: {'completed': 0, 'cancelled': 0, 'total': 0, 'gmv': 0, 'settlement': 0})
for r in completed:
    order_time = r[4]
    if order_time is None:
        continue
    if isinstance(order_time, datetime):
        date_key = order_time.strftime('%Y-%m-%d')
    else:
        continue
    daily_orders[date_key]['completed'] += 1
    daily_orders[date_key]['total'] += 1
    daily_orders[date_key]['gmv'] += r[3] if r[3] is not None else 0
    daily_orders[date_key]['settlement'] += r[7] if r[7] is not None else 0

for r in cancelled:
    order_time = r[4]
    if order_time is None:
        continue
    if isinstance(order_time, datetime):
        date_key = order_time.strftime('%Y-%m-%d')
    else:
        continue
    daily_orders[date_key]['cancelled'] += 1
    daily_orders[date_key]['total'] += 1

print(f"{'日期':<14} {'总订单':<8} {'已完成':<8} {'已取消':<8} {'取消率':<8} {'GMV':<14} {'结算金额':<14}")
print("-" * 74)
for date_key in sorted(daily_orders.keys()):
    d = daily_orders[date_key]
    cancel_rate = d['cancelled'] / d['total'] * 100 if d['total'] > 0 else 0
    print(f"{date_key:<14} {d['total']:<8} {d['completed']:<8} {d['cancelled']:<8} {cancel_rate:<7.1f}% ₱{d['gmv']:<10,.2f} ₱{d['settlement']:<10,.2f}")

print()

# ========== 4. 取消订单分析 ==========
print("=" * 70)
print("四、取消订单分析")
print("=" * 70)

cancel_by_who = Counter()
cancel_gmv_total = 0
cancel_with_refund = 0
cancel_refund_total = 0

for r in cancelled:
    who = r[6]  # 取消方
    if who is not None:
        cancel_by_who[who] += 1
    gmv = r[3] if r[3] is not None else 0
    cancel_gmv_total += gmv
    refund = r[2] if r[2] is not None else 0
    if refund > 0:
        cancel_with_refund += 1
        cancel_refund_total += refund

print(f"取消订单总数: {len(cancelled)}")
print(f"取消订单GMV总额: ₱{cancel_gmv_total:,.2f}")
print(f"有退款的取消订单: {cancel_with_refund} 单")
print(f"取消退款总额: ₱{cancel_refund_total:,.2f}")
print()
print("取消方分布:")
for who, count in cancel_by_who.most_common():
    pct = count / len(cancelled) * 100
    print(f"  {who}: {count} 单 ({pct:.1f}%)")

print()

# ========== 5. 商品价格分布 ==========
print("=" * 70)
print("五、商品价格分布分析")
print("=" * 70)

price_ranges = defaultdict(lambda: {'completed': 0, 'cancelled': 0, 'settlement': 0})
for r in completed:
    price = r[3] if r[3] is not None else 0
    if price <= 0:
        continue
    if price <= 80:
        bucket = "≤ ₱80"
    elif price <= 100:
        bucket = "₱81-100"
    elif price <= 130:
        bucket = "₱101-130"
    elif price <= 150:
        bucket = "₱131-150"
    elif price <= 200:
        bucket = "₱151-200"
    else:
        bucket = "> ₱200"
    price_ranges[bucket]['completed'] += 1
    price_ranges[bucket]['settlement'] += r[7] if r[7] is not None else 0

for r in cancelled:
    price = r[3] if r[3] is not None else 0
    if price <= 0:
        continue
    if price <= 80:
        bucket = "≤ ₱80"
    elif price <= 100:
        bucket = "₱81-100"
    elif price <= 130:
        bucket = "₱101-130"
    elif price <= 150:
        bucket = "₱131-150"
    elif price <= 200:
        bucket = "₱151-200"
    else:
        bucket = "> ₱200"
    price_ranges[bucket]['cancelled'] += 1

print(f"{'价格区间':<14} {'已完成':<8} {'已取消':<8} {'总订单':<8} {'取消率':<8} {'结算金额':<14}")
print("-" * 60)
for bucket in ["≤ ₱80", "₱81-100", "₱101-130", "₱131-150", "₱151-200", "> ₱200"]:
    d = price_ranges[bucket]
    total = d['completed'] + d['cancelled']
    if total == 0:
        continue
    cancel_rate = d['cancelled'] / total * 100
    print(f"{bucket:<14} {d['completed']:<8} {d['cancelled']:<8} {total:<8} {cancel_rate:<7.1f}% ₱{d['settlement']:<10,.2f}")

print()

# ========== 6. 费用分析 ==========
print("=" * 70)
print("六、费用分析")
print("=" * 70)

total_gmv = sum(r[3] for r in completed if r[3] is not None)
total_platform_fee = sum(abs(r[9]) for r in completed if r[9] is not None and isinstance(r[9], (int, float)))
total_commission = sum(abs(r[10]) for r in completed if r[10] is not None and isinstance(r[10], (int, float)))
total_shipping = sum(abs(r[11]) for r in completed if r[11] is not None and isinstance(r[11], (int, float)))
total_coupon = sum(abs(r[12]) for r in completed if r[12] is not None and isinstance(r[12], (int, float)))
total_affiliate = sum(abs(r[13]) for r in completed if r[13] is not None and isinstance(r[13], (int, float)))
total_other = sum(abs(r[14]) for r in completed if r[14] is not None and isinstance(r[14], (int, float)))
total_other2 = sum(abs(r[15]) for r in completed if r[15] is not None and isinstance(r[15], (int, float)))

total_fees = total_platform_fee + total_commission + total_shipping + total_coupon + total_affiliate + total_other + total_other2

print(f"{'费用项目':<16} {'金额(₱)':<14} {'占GMV比例':<12}")
print("-" * 42)
print(f"{'商品总额(GMV)':<16} ₱{total_gmv:<10,.2f} {'100%':<12}")
print(f"{'平台费用':<16} ₱{total_platform_fee:<10,.2f} {total_platform_fee/total_gmv*100 if total_gmv else 0:<11.2f}%")
print(f"{'佣金':<16} ₱{total_commission:<10,.2f} {total_commission/total_gmv*100 if total_gmv else 0:<11.2f}%")
print(f"{'运费':<16} ₱{total_shipping:<10,.2f} {total_shipping/total_gmv*100 if total_gmv else 0:<11.2f}%")
print(f"{'平台优惠券':<16} ₱{total_coupon:<10,.2f} {total_coupon/total_gmv*100 if total_gmv else 0:<11.2f}%")
print(f"{'达人佣金':<16} ₱{total_affiliate:<10,.2f} {total_affiliate/total_gmv*100 if total_gmv else 0:<11.2f}%")
print(f"{'其他扣费1':<16} ₱{total_other:<10,.2f} {total_other/total_gmv*100 if total_gmv else 0:<11.2f}%")
print(f"{'其他扣费2':<16} ₱{total_other2:<10,.2f} {total_other2/total_gmv*100 if total_gmv else 0:<11.2f}%")
print("-" * 42)
print(f"{'总费用':<16} ₱{total_fees:<10,.2f} {total_fees/total_gmv*100 if total_gmv else 0:<11.2f}%")
print(f"{'净结算金额':<16} ₱{total_settlement:<10,.2f} {total_settlement/total_gmv*100 if total_gmv else 0:<11.2f}%")
print()

# ========== 7. 结算周期分析 ==========
print("=" * 70)
print("七、结算周期分析（下单到结算的天数）")
print("=" * 70)

settle_days = []
for r in completed:
    order_time = r[4]
    settle_time = r[5]
    if order_time is not None and settle_time is not None:
        if isinstance(order_time, datetime) and isinstance(settle_time, datetime):
            days = (settle_time - order_time).days
            settle_days.append(days)

if settle_days:
    avg_days = sum(settle_days) / len(settle_days)
    min_days = min(settle_days)
    max_days = max(settle_days)
    print(f"平均结算周期: {avg_days:.1f} 天")
    print(f"最短结算周期: {min_days} 天")
    print(f"最长结算周期: {max_days} 天")
    print()
    
    # 分布
    day_dist = Counter(settle_days)
    print("结算周期分布:")
    print(f"{'天数':<8} {'订单数':<8} {'占比':<8}")
    print("-" * 24)
    for days in sorted(day_dist.keys()):
        pct = day_dist[days] / len(settle_days) * 100
        bar = '█' * int(pct / 2)
        print(f"{days:<8} {day_dist[days]:<8} {pct:<7.1f}% {bar}")

print()

# ========== 8. 退款分析 ==========
print("=" * 70)
print("八、退款分析")
print("=" * 70)

refund_orders = [r for r in completed if r[2] is not None and r[2] > 0]
total_refund = sum(r[2] for r in refund_orders)
print(f"有退款的已完成订单: {len(refund_orders)} 单")
print(f"退款总额: ₱{total_refund:,.2f}")
print(f"退款率(按订单数): {len(refund_orders)/len(completed)*100:.2f}%")
print(f"退款率(按金额): {total_refund/total_gmv*100:.2f}%")
print()

# ========== 9. 达人佣金分析 ==========
print("=" * 70)
print("九、达人佣金分析")
print("=" * 70)

affiliate_orders = [r for r in completed if r[13] is not None and isinstance(r[13], (int, float)) and r[13] < 0]
no_affiliate_orders = [r for r in completed if r[13] is None or (isinstance(r[13], (int, float)) and r[13] == 0)]

print(f"有达人佣金的订单: {len(affiliate_orders)} 单")
print(f"无达人佣金的订单: {len(no_affiliate_orders)} 单")
print(f"达人佣金总支出: ₱{total_affiliate:,.2f}")
print(f"达人佣金订单占比: {len(affiliate_orders)/len(completed)*100:.1f}%")
print()

# 有达人佣金 vs 无达人佣金的平均结算金额
if affiliate_orders:
    avg_affiliate_settle = sum(r[7] for r in affiliate_orders if r[7] is not None) / len(affiliate_orders)
    print(f"有达人佣金订单平均结算: ₱{avg_affiliate_settle:,.2f}")
if no_affiliate_orders:
    avg_no_affiliate_settle = sum(r[7] for r in no_affiliate_orders if r[7] is not None) / len(no_affiliate_orders)
    print(f"无达人佣金订单平均结算: ₱{avg_no_affiliate_settle:,.2f}")

print()

# ========== 10. 周度趋势 ==========
print("=" * 70)
print("十、周度趋势")
print("=" * 70)

weekly_data = defaultdict(lambda: {'gmv': 0, 'settlement': 0, 'count': 0, 'cancelled': 0})
for r in completed:
    order_time = r[4]
    if order_time is None or not isinstance(order_time, datetime):
        continue
    week_start = order_time - timedelta(days=order_time.weekday())
    week_key = week_start.strftime('%Y-%m-%d')
    weekly_data[week_key]['gmv'] += r[3] if r[3] is not None else 0
    weekly_data[week_key]['settlement'] += r[7] if r[7] is not None else 0
    weekly_data[week_key]['count'] += 1

for r in cancelled:
    order_time = r[4]
    if order_time is None or not isinstance(order_time, datetime):
        continue
    week_start = order_time - timedelta(days=order_time.weekday())
    week_key = week_start.strftime('%Y-%m-%d')
    weekly_data[week_key]['cancelled'] += 1

print(f"{'周起始':<14} {'订单数':<8} {'取消数':<8} {'GMV':<14} {'结算金额':<14}")
print("-" * 58)
for week_key in sorted(weekly_data.keys()):
    d = weekly_data[week_key]
    print(f"{week_key:<14} {d['count']:<8} {d['cancelled']:<8} ₱{d['gmv']:<10,.2f} ₱{d['settlement']:<10,.2f}")

print()

# ========== 11. 关键指标汇总 ==========
print("=" * 70)
print("十一、关键指标汇总")
print("=" * 70)

# 净结算率
net_rate = total_settlement / total_gmv * 100 if total_gmv else 0
# 客单价
avg_order_value = total_gmv / len(completed) if completed else 0
# 平均每单净收入
avg_net_per_order = total_settlement / len(completed) if completed else 0

print(f"总 GMV: ₱{total_gmv:,.2f}")
print(f"总结算金额: ₱{total_settlement:,.2f}")
print(f"净结算率: {net_rate:.2f}%")
print(f"客单价: ₱{avg_order_value:,.2f}")
print(f"平均每单净收入: ₱{avg_net_per_order:,.2f}")
print(f"总费用: ₱{total_fees:,.2f} (占GMV的 {total_fees/total_gmv*100:.2f}%)")
print(f"取消订单数: {len(cancelled)}")
print(f"取消率: {len(cancelled)/(len(completed)+len(cancelled))*100:.2f}%")
print(f"退款订单数: {len(refund_orders)}")
print(f"退款率: {len(refund_orders)/len(completed)*100:.2f}%")
