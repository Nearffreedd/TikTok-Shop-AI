"""
联盟销售数据导入脚本
功能：读取 AllPlan 导出的 ListProducts Excel → 写入 SQLite 数据库 affiliate_weekly 表
用法：python import_affiliate_data.py <Excel文件路径> [店铺代码]

示例：
  python import_affiliate_data.py "Layer2_Working/ListProducts_2026-05-11-2026-05-18_ALLPlan_20260526083006.xlsx" PNB
  python import_affiliate_data.py "Layer2_Working/ListProducts_2026-05-18-2026-05-25_ALLPlan_20260526083103.xlsx"
  （不指定店铺代码时，默认 PNB）
"""

import sqlite3
import openpyxl
import os
import sys
import re
from datetime import datetime

# 数据库路径
DB_PATH = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                       'Layer1_Permanent', '01_My_Shop', 'shop_data.db')


def parse_week_range_from_filename(filename):
    """从文件名解析周范围
    格式: 'ListProducts_2026-05-11-2026-05-18_ALLPlan_...' -> ('2026-05-11', '2026-05-18')
    """
    match = re.search(r'(\d{4}-\d{2}-\d{2})-(\d{4}-\d{2}-\d{2})', filename)
    if match:
        return match.group(1), match.group(2)
    return None, None


def parse_float(val):
    """安全解析数值，处理 ₱ 符号和千位分隔符"""
    if val is None or val == '' or val == '-' or val == '--':
        return 0.0
    try:
        s = str(val).strip()
        s = s.replace('PHP', '').replace('₱', '').replace('$', '').replace(',', '').strip()
        return float(s)
    except (ValueError, TypeError):
        return 0.0


def parse_int(val):
    """安全解析整数"""
    if val is None or val == '' or val == '-' or val == '--':
        return 0
    try:
        return int(float(str(val).replace(',', '').strip()))
    except (ValueError, TypeError):
        return 0


def import_affiliate_excel(filepath, shop_code='PNB'):
    """将联盟销售数据 Excel 导入数据库"""

    if not os.path.exists(filepath):
        print(f'[X] 文件不存在: {filepath}')
        return False

    filename = os.path.basename(filepath)

    # 从文件名解析周范围
    week_start, week_end = parse_week_range_from_filename(filename)
    if week_start is None:
        print(f'[X] 无法从文件名解析日期范围: {filename}')
        print(f'   期望格式: ListProducts_YYYY-MM-DD-YYYY-MM-DD_...')
        return False

    # 读取 Excel
    wb = openpyxl.load_workbook(filepath)
    ws = wb.active

    print(f'[文件] {filename}')
    print(f'[店铺] {shop_code}')
    print(f'[周期] {week_start} ~ {week_end}')
    print(f'[数据] {ws.max_row - 1} 条商品记录')

    # 读取表头（第1行）
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]

    # 连接数据库
    conn = sqlite3.connect(DB_PATH)
    cursor = conn.cursor()

    # 获取店铺ID
    cursor.execute('SELECT shop_id FROM shops WHERE shop_code = ?', (shop_code,))
    row = cursor.fetchone()
    if row is None:
        print(f'[X] 店铺代码 {shop_code} 不存在，请先初始化数据库')
        conn.close()
        return False
    shop_id = row[0]

    imported_count = 0
    skipped_count = 0
    new_product_count = 0

    # 遍历数据行（从第2行开始）
    for row_data in ws.iter_rows(min_row=2, values_only=True):
        if row_data[0] is None:
            continue

        d = dict(zip(headers, row_data))

        platform_product_id = str(d.get('商品 ID', '')).strip()
        product_name = str(d.get('商品名称', '')).strip()

        if not platform_product_id or platform_product_id == 'None':
            continue

        # 解析数值
        affiliate_gmv = parse_float(d.get('GMV'))
        affiliate_units = parse_int(d.get('成交件数'))
        estimated_commission = parse_float(d.get('预计佣金'))
        sample_count = parse_int(d.get('样品数'))
        selling_creators = parse_int(d.get('销售达人'))
        live_count = parse_int(d.get('直播数'))
        video_count = parse_int(d.get('视频数量'))
        refunded_gmv = parse_float(d.get('已退款的 GMV'))
        refunded_units = parse_int(d.get('已退款的成交件数'))
        fixed_fee = parse_float(d.get('预计固定费用'))

        # 1. 插入或更新商品（确保 products 表中有此商品）
        cursor.execute('''
            INSERT INTO products (shop_id, platform_product_id, product_name, status, first_seen_date, last_seen_date)
            VALUES (?, ?, ?, 'Active', ?, ?)
            ON CONFLICT(shop_id, platform_product_id) DO UPDATE SET
                product_name = excluded.product_name,
                last_seen_date = excluded.last_seen_date,
                updated_at = CURRENT_TIMESTAMP
        ''', (shop_id, platform_product_id, product_name, week_start, week_end))

        # 获取 product_id
        cursor.execute('SELECT product_id FROM products WHERE shop_id = ? AND platform_product_id = ?',
                      (shop_id, platform_product_id))
        product_id = cursor.fetchone()[0]

        # 2. 插入联盟周数据
        insert_sql = """
            INSERT INTO affiliate_weekly (
                product_id, week_start, week_end,
                affiliate_gmv, affiliate_units, estimated_commission,
                sample_count, selling_creators, live_count, video_count,
                refunded_gmv, refunded_units, fixed_fee,
                source_file
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """
        insert_values = (
            product_id, week_start, week_end,
            affiliate_gmv, affiliate_units, estimated_commission,
            sample_count, selling_creators, live_count, video_count,
            refunded_gmv, refunded_units, fixed_fee,
            filename
        )

        try:
            cursor.execute(insert_sql, insert_values)
            imported_count += 1
        except sqlite3.IntegrityError:
            # 已存在则更新
            update_sql = """
                UPDATE affiliate_weekly SET
                    affiliate_gmv = ?, affiliate_units = ?, estimated_commission = ?,
                    sample_count = ?, selling_creators = ?, live_count = ?, video_count = ?,
                    refunded_gmv = ?, refunded_units = ?, fixed_fee = ?,
                    source_file = ?
                WHERE product_id = ? AND week_start = ? AND week_end = ?
            """
            update_values = (
                affiliate_gmv, affiliate_units, estimated_commission,
                sample_count, selling_creators, live_count, video_count,
                refunded_gmv, refunded_units, fixed_fee,
                filename,
                product_id, week_start, week_end
            )
            cursor.execute(update_sql, update_values)
            skipped_count += 1

    conn.commit()
    conn.close()

    print(f'[完成] 联盟数据导入完成: 新增 {imported_count} 条, 更新 {skipped_count} 条')
    return True


if __name__ == '__main__':
    if len(sys.argv) < 2:
        print('用法: python import_affiliate_data.py <Excel文件路径> [店铺代码]')
        print('示例: python import_affiliate_data.py "Layer2_Working/ListProducts_2026-05-11-2026-05-18_ALLPlan_20260526083006.xlsx" PNB')
        sys.exit(1)

    filepath = sys.argv[1]
    shop_code = sys.argv[2] if len(sys.argv) > 2 else 'PNB'

    # 如果是相对路径，拼接
    if not os.path.isabs(filepath):
        filepath = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), filepath)

    import_affiliate_excel(filepath, shop_code)
