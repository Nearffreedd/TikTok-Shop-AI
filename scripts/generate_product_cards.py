# -*- coding: utf-8 -*-
"""
从 星昊科技有限公司SKU生成.xlsx 读取数据，
按款式代码（类别编码+产品编码）分组，
生成 Obsidian Markdown 产品信息卡 + 产品款式总表。
"""
import openpyxl, os, json
from collections import defaultdict

# ========== 配置 ==========
EXCEL_PATH = 'Layer2_Working/星昊科技有限公司SKU生成.xlsx'
OUTPUT_DIR = 'Layer1_Permanent/02_Products/产品信息卡'
INDEX_FILE = 'Layer1_Permanent/02_Products/产品款式总表.md'

# ========== 读取数据 ==========
wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
ws = wb['单品SKU生成表']

# 按款式代码（类别编码+产品编码）分组
groups = defaultdict(list)
for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True):
    vals = [str(v) if v is not None else '' for v in row]
    ov_sku = vals[27].strip() if len(vals) > 27 else ''
    if not ov_sku:
        continue
    
    cat_code = vals[8].strip()    # 类别编码
    prod_code = vals[9].strip()   # 产品编码
    style_code = vals[10].strip() # 风格编码
    style_key = f'{cat_code}{prod_code}'  # 款式代码，如 FQ01
    
    groups[style_key].append({
        'ov_sku': ov_sku,
        'style_code': style_code,
        'color': vals[6].strip(),
        'size': vals[4].strip(),
        'mat': vals[3].strip(),
        'name': vals[16].strip(),
        'img': vals[17].strip(),
        'weight': vals[33].strip() if len(vals) > 33 else '',
        'ref_price': vals[29].strip() if len(vals) > 29 else '',
        'cost': vals[18].strip() if len(vals) > 18 else '',
        'prod_name': vals[0].strip(),  # 产品类别中文名
        'prod_en': vals[21].strip() if len(vals) > 21 else '',
        'color_en': vals[26].strip() if len(vals) > 26 else '',
        'size_en': vals[24].strip() if len(vals) > 24 else '',
        'length': vals[30].strip() if len(vals) > 30 else '',
        'width': vals[31].strip() if len(vals) > 31 else '',
        'height': vals[32].strip() if len(vals) > 32 else '',
    })

print(f'共读取到 {sum(len(v) for v in groups.values())} 个SKU，{len(groups)} 个款式')

# ========== 创建输出目录 ==========
os.makedirs(OUTPUT_DIR, exist_ok=True)

# ========== 生成每个款式的信息卡 ==========
index_rows = []  # 用于总表的数据

for style_key in sorted(groups.keys()):
    items = groups[style_key]
    
    # 提取款式信息
    prod_name = items[0]['prod_name']
    mat = items[0]['mat']
    colors = sorted(set(i['color'] for i in items))
    sizes = sorted(set(i['size'] for i in items))
    styles = sorted(set(i['style_code'] for i in items))
    first_img = items[0]['img']
    ref_price = items[0]['ref_price']
    cost = items[0]['cost']
    
    # 构建 YAML Front Matter
    yaml_lines = ['---']
    yaml_lines.append(f'款式代码: {style_key}')
    yaml_lines.append(f'产品名称: {prod_name}')
    yaml_lines.append(f'类目: {prod_name}')
    yaml_lines.append(f'材质: {mat}')
    yaml_lines.append(f'颜色列表:')
    for c in colors:
        yaml_lines.append(f'  - {c}')
    yaml_lines.append(f'规格尺寸: {", ".join(sizes) if sizes else ""}')
    yaml_lines.append(f'重量(g): {items[0]["weight"]}')
    yaml_lines.append(f'产品图: {first_img}')
    yaml_lines.append(f'参考价(RMB): {ref_price}')
    yaml_lines.append(f'成本(RMB): {cost}')
    yaml_lines.append(f'适用人群: 待补充')
    yaml_lines.append(f'卖点: 待补充')
    yaml_lines.append(f'上架日期: 待补充')
    yaml_lines.append(f'海外仓SKU列表:')
    for item in items:
        yaml_lines.append(f'  - {item["ov_sku"]}')
    yaml_lines.append('---')
    
    # 构建 Markdown 正文
    md_lines = []
    md_lines.append(f'## 款式总览')
    md_lines.append('')
    md_lines.append('| 字段 | 值 |')
    md_lines.append('|------|-----|')
    md_lines.append(f'| 款式代码 | {style_key} |')
    md_lines.append(f'| 产品类别 | {prod_name} |')
    md_lines.append(f'| 材质 | {mat} |')
    md_lines.append(f'| 规格尺寸 | {", ".join(sizes) if sizes else "-"} |')
    md_lines.append(f'| 颜色 | {", ".join(colors) if colors else "-"} |')
    md_lines.append(f'| 风格 | {", ".join(styles) if styles else "-"} |')
    md_lines.append(f'| 参考价(RMB) | {ref_price if ref_price else "-"} |')
    md_lines.append(f'| 成本(RMB) | {cost if cost else "-"} |')
    md_lines.append(f'| 重量(g) | {items[0]["weight"] if items[0]["weight"] else "-"} |')
    md_lines.append(f'| 适用人群 | 待补充 |')
    md_lines.append(f'| 卖点 | 待补充 |')
    md_lines.append(f'| 上架日期 | 待补充 |')
    md_lines.append('')
    
    # 产品图
    if first_img:
        md_lines.append(f'![产品图]({first_img})')
        md_lines.append('')
    
    # SKU 清单
    md_lines.append(f'## SKU 清单')
    md_lines.append('')
    md_lines.append(f'共 **{len(items)}** 个SKU')
    md_lines.append('')
    md_lines.append('| 海外仓SKU | 颜色 | 尺寸 | 风格 | 重量(g) | 产品图 |')
    md_lines.append('|-----------|------|------|------|---------|--------|')
    for item in items:
        img_link = f'[图]({item["img"]})' if item['img'] else '-'
        md_lines.append(f'| {item["ov_sku"]} | {item["color"]} | {item["size"]} | {item["style_code"]} | {item["weight"] if item["weight"] else "-"} | {img_link} |')
    md_lines.append('')
    
    # 完整内容
    content = '\n'.join(yaml_lines) + '\n' + '\n'.join(md_lines)
    
    # 写入文件
    filepath = os.path.join(OUTPUT_DIR, f'{style_key}.md')
    with open(filepath, 'w', encoding='utf-8') as f:
        f.write(content)
    
    # 收集总表数据
    index_rows.append({
        '款式代码': style_key,
        '产品类别': prod_name,
        '材质': mat,
        'SKU数量': len(items),
        '颜色': ', '.join(colors),
        '风格': ', '.join(styles),
        '尺寸': ', '.join(sizes),
        '参考价': ref_price if ref_price else '-',
        '成本': cost if cost else '-',
        '重量': items[0]['weight'] if items[0]['weight'] else '-',
        '文件路径': f'产品信息卡/{style_key}.md',
    })
    
    print(f'  ✓ {style_key}.md ({prod_name}, {len(items)}个SKU)')

# ========== 生成产品款式总表 ==========
print(f'\n生成产品款式总表...')

total_lines = []
total_lines.append('---')
total_lines.append('创建时间: 自动生成')
total_lines.append('更新说明: 从星昊科技有限公司SKU生成.xlsx 自动生成')
total_lines.append('---')
total_lines.append('')
total_lines.append('# 产品款式总表')
total_lines.append('')
total_lines.append(f'> 共 **{len(index_rows)}** 个款式，**{sum(v["SKU数量"] for v in index_rows)}** 个SKU')
total_lines.append('')
total_lines.append('| 款式代码 | 产品类别 | 材质 | SKU数 | 颜色 | 风格 | 尺寸 | 参考价 | 成本 | 信息卡 |')
total_lines.append('|---------|---------|------|------|------|------|------|-------|------|-------|')

for r in index_rows:
    total_lines.append(f'| {r["款式代码"]} | {r["产品类别"]} | {r["材质"]} | {r["SKU数量"]} | {r["颜色"]} | {r["风格"]} | {r["尺寸"]} | {r["参考价"]} | {r["成本"]} | [{r["款式代码"]}]({r["文件路径"]}) |')

total_lines.append('')
total_lines.append('---')
total_lines.append('')
total_lines.append('## 按类目统计')
total_lines.append('')

# 按类目统计
from collections import Counter
cat_counter = Counter(r['产品类别'] for r in index_rows)
total_lines.append('| 类目 | 款式数 |')
total_lines.append('|------|-------|')
for cat, count in sorted(cat_counter.items()):
    total_lines.append(f'| {cat} | {count} |')

with open(INDEX_FILE, 'w', encoding='utf-8') as f:
    f.write('\n'.join(total_lines))

print(f'  ✓ 产品款式总表.md 已生成')
print(f'\n✅ 全部完成！')
print(f'  - 产品信息卡目录: {OUTPUT_DIR}/')
print(f'  - 款式总表: {INDEX_FILE}')
