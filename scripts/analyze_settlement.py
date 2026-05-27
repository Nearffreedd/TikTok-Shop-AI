import openpyxl
import sys
import os

filepath = os.path.join(os.path.dirname(os.path.dirname(__file__)), 'Layer2_Working', '结算统计.xlsx')
print(f"File path: {filepath}")
print(f"File exists: {os.path.exists(filepath)}")

wb = openpyxl.load_workbook(filepath, data_only=True)
print(f'Sheet names: {wb.sheetnames}')

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    print(f'\n{"="*80}')
    print(f'=== Sheet: {sheet_name} ===')
    print(f'Rows: {ws.max_row}, Cols: {ws.max_column}')
    print(f'{"="*80}')
    
    # Print ALL rows
    for row_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True), 1):
        print(f'Row {row_idx}: {list(row)}')
