"""
每周商品数据导入脚本
功能：读取 Excel -> 写入 SQLite 数据库
用法：python import_weekly_data.py <Excel文件路径> [店铺代码]

示例：
  python import_weekly_data.py "Layer2_Working/PNB-商品数据 20260518-20260524.xlsx" PNB
  python import_weekly_data.py "Layer2_Working/PNB-商品数据 20260518-20260524.xlsx"
  （不指定店铺代码时，从文件名自动识别 PNB）
"""

import sqlite3
import openpyxl
import os
import sys
import re
from datetime import datetime

# 数据库路径
DB_PATH = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                       'Layer1_Permanent', '01_My_Shop', 'shop_data.db')


def parse_week_range_from_filename(filename):
    """从文件名解析周范围，如 'PNB-商品数据 20260518-20260524.xlsx' -> ('2026-05-18', '2026-05-24')"""
    match = re.search(r'(\d{8})-(\d{8})', filename)
    if match:
        start = datetime.strptime(match.group(1), '%Y%m%d').strftime('%Y-%m-%d')
        end = datetime.strptime(match.group(2), '%Y%m%d').strftime('%Y-%m-%d')
        return start, end
    return None, None


def parse_week_range_from_a1(cell_value):
    """从 Excel A1 单元格解析日期范围
    格式: '2026-05-18 ~ 2026-05-24' -> ('2026-05-18', '2026-05-24')
    """
    if cell_value is None:
        return None, None
    text = str(cell_value).strip()
    match = re.search(r'(\d{4}-\d{2}-\d{2})\s*~\s*(\d{4}-\d{2}-\d{2})', text)
    if match:
        return match.group(1), match.group(2)
    return None, None


def detect_shop_code(filename):
    """从文件名识别店铺代码，如 'PNB-商品数据...' -> 'PNB'"""
    match = re.match(r'([A-Z]+)-商品数据', filename)
    if match:
        return match.group(1)
    return None


def parse_float(val):
    """安全解析数值，处理 ₱ 符号和千位分隔符"""
    if val is None or val == '' or val == '-':
        return 0
    try:
        s = str(val).strip()
        # 移除货币符号（₱、PHP、$ 等）和千位分隔符
        s = s.replace('PHP', '').replace('₱', '').replace('$', '').replace(',', '').strip()
        return float(s)
    except:
        return 0


def parse_int(val):
    """安全解析整数"""
    if val is None or val == '' or val == '-':
        return 0
    try:
        return int(float(str(val).replace(',', '').strip()))
    except:
        return 0


def parse_pct(val):
    """安全解析百分比"""
    if val is None or val == '' or val == '-':
        return 0.0
    try:
        return float(str(val).replace('%', '').strip()) / 100
    except:
        return 0.0


def import_excel_to_db(filepath, shop_code=None):
    """将Excel商品数据导入数据库"""

    if not os.path.exists(filepath):
        print(f'[X] 文件不存在: {filepath}')
        return False

    filename = os.path.basename(filepath)

    # 自动识别店铺代码
    if shop_code is None:
        shop_code = detect_shop_code(filename)
        if shop_code is None:
            print('[X] 无法从文件名识别店铺代码，请手动指定')
            return False

    # 读取Excel
    wb = openpyxl.load_workbook(filepath)
    ws = wb.active

    # 优先从 A1 单元格读取日期周期
    a1_value = ws.cell(1, 1).value
    week_start, week_end = parse_week_range_from_a1(a1_value)

    # 如果 A1 没有，回退到从文件名解析
    if week_start is None:
        week_start, week_end = parse_week_range_from_filename(filename)

    if week_start is None:
        print('[X] 无法解析日期范围（A1单元格和文件名均无效）')
        return False

    print(f'[文件] {filename}')
    print(f'[店铺] {shop_code}')
    print(f'[周期] {week_start} ~ {week_end}')

    # 读取表头（第3行）
    headers = [c.value for c in next(ws.iter_rows(min_row=3, max_row=3))]

    # 连接数据库
    conn = sqlite3.connect(DB_PATH)
    cursor = conn.cursor()

    # 获取店铺ID
    cursor.execute('SELECT shop_id FROM shops WHERE shop_code = ?', (shop_code,))
    row = cursor.fetchone()
    if row is None:
        print(f'[X] 店铺代码 {shop_code} 不存在，请先初始化数据库')
        conn.close()
        return False
    shop_id = row[0]

    imported_count = 0
    skipped_count = 0

    # 遍历数据行（从第4行开始）
    for row_data in ws.iter_rows(min_row=4, values_only=True):
        if row_data[0] is None:
            continue

        d = dict(zip(headers, row_data))

        platform_product_id = str(d.get('ID', '')).strip()
        product_name = str(d.get('商品', '')).strip()
        status = str(d.get('状态', 'Active')).strip()

        if not platform_product_id or platform_product_id == 'None':
            continue

        # 解析数值
        gmv = parse_float(d.get('GMV'))
        units = parse_int(d.get('成交件数'))
        orders = parse_int(d.get('订单数'))
        refund_units = parse_float(d.get('已退款的商品件数'))

        # 商城渠道
        shop_gmv = parse_float(d.get('商城页 GMV'))
        shop_units = parse_int(d.get('商城商品成交件数'))
        shop_imp = parse_int(d.get('商城页发品曝光次数'))
        shop_pv = parse_int(d.get('商城页面浏览次数'))
        shop_uv = parse_int(d.get('商城页去重商品客户数'))
        shop_ctr = parse_pct(d.get('商城点击率'))
        shop_cvr = parse_pct(d.get('商城转化率'))

        # 视频渠道
        video_gmv = parse_float(d.get('视频归因 GMV'))
        video_units = parse_int(d.get('视频归因成交件数'))
        video_imp = parse_int(d.get('视频曝光次数'))
        video_pv = parse_int(d.get('来自视频的页面浏览次数'))
        video_uv = parse_int(d.get('视频去重商品客户数'))
        video_ctr = parse_pct(d.get('视频点击率'))
        video_cvr = parse_pct(d.get('视频转化率'))

        # 直播渠道
        live_gmv = parse_float(d.get('直播归因 GMV'))
        live_units = parse_int(d.get('直播归因成交件数'))
        live_imp = parse_int(d.get('直播曝光次数'))
        live_pv = parse_int(d.get('直播的页面浏览次数'))
        live_uv = parse_int(d.get('直播去重商品客户数'))
        live_ctr = parse_pct(d.get('直播点击率'))
        live_cvr = parse_pct(d.get('直播转化率'))

        # 商品卡渠道
        card_gmv = parse_float(d.get('商品卡归因 GMV'))
        card_units = parse_int(d.get('商品卡归因成交件数'))
        card_imp = parse_int(d.get('商品卡曝光次数'))
        card_pv = parse_int(d.get('商品卡的页面浏览次数'))
        card_uv = parse_int(d.get('商品卡去重客户数'))
        card_ctr = parse_pct(d.get('商品卡点击率'))
        card_cvr = parse_pct(d.get('商品卡转化率'))

        # 1. 插入或更新商品
        cursor.execute('''
            INSERT INTO products (shop_id, platform_product_id, product_name, status, first_seen_date, last_seen_date)
            VALUES (?, ?, ?, ?, ?, ?)
            ON CONFLICT(shop_id, platform_product_id) DO UPDATE SET
                product_name = excluded.product_name,
                status = excluded.status,
                last_seen_date = excluded.last_seen_date,
                updated_at = CURRENT_TIMESTAMP
        ''', (shop_id, platform_product_id, product_name, status, week_start, week_end))

        # 获取 product_id
        cursor.execute('SELECT product_id FROM products WHERE shop_id = ? AND platform_product_id = ?',
                      (shop_id, platform_product_id))
        product_id = cursor.fetchone()[0]

        # 2. 插入周数据
        insert_sql = """
            INSERT INTO product_weekly (
                product_id, week_start, week_end,
                gmv, units_sold, orders, refund_units,
                shop_gmv, shop_units, shop_impressions, shop_page_views, shop_unique_visitors,
                shop_click_rate, shop_conversion_rate,
                video_gmv, video_units, video_impressions, video_page_views, video_unique_visitors,
                video_click_rate, video_conversion_rate,
                live_gmv, live_units, live_impressions, live_page_views, live_unique_visitors,
                live_click_rate, live_conversion_rate,
                card_gmv, card_units, card_impressions, card_page_views, card_unique_visitors,
                card_click_rate, card_conversion_rate,
                source_file
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """
        insert_values = (
            product_id, week_start, week_end,
            gmv, units, orders, refund_units,
            shop_gmv, shop_units, shop_imp, shop_pv, shop_uv,
            shop_ctr, shop_cvr,
            video_gmv, video_units, video_imp, video_pv, video_uv,
            video_ctr, video_cvr,
            live_gmv, live_units, live_imp, live_pv, live_uv,
            live_ctr, live_cvr,
            card_gmv, card_units, card_imp, card_pv, card_uv,
            card_ctr, card_cvr,
            filename
        )
        try:
            cursor.execute(insert_sql, insert_values)
            imported_count += 1
        except sqlite3.IntegrityError:
            # 已存在则更新
            update_sql = """
                UPDATE product_weekly SET
                    gmv = ?, units_sold = ?, orders = ?, refund_units = ?,
                    shop_gmv = ?, shop_units = ?, shop_impressions = ?, shop_page_views = ?, shop_unique_visitors = ?,
                    shop_click_rate = ?, shop_conversion_rate = ?,
                    video_gmv = ?, video_units = ?, video_impressions = ?, video_page_views = ?, video_unique_visitors = ?,
                    video_click_rate = ?, video_conversion_rate = ?,
                    live_gmv = ?, live_units = ?, live_impressions = ?, live_page_views = ?, live_unique_visitors = ?,
                    live_click_rate = ?, live_conversion_rate = ?,
                    card_gmv = ?, card_units = ?, card_impressions = ?, card_page_views = ?, card_unique_visitors = ?,
                    card_click_rate = ?, card_conversion_rate = ?,
                    source_file = ?
                WHERE product_id = ? AND week_start = ? AND week_end = ?
            """
            update_values = (
                gmv, units, orders, refund_units,
                shop_gmv, shop_units, shop_imp, shop_pv, shop_uv,
                shop_ctr, shop_cvr,
                video_gmv, video_units, video_imp, video_pv, video_uv,
                video_ctr, video_cvr,
                live_gmv, live_units, live_imp, live_pv, live_uv,
                live_ctr, live_cvr,
                card_gmv, card_units, card_imp, card_pv, card_uv,
                card_ctr, card_cvr,
                filename,
                product_id, week_start, week_end
            )
            cursor.execute(update_sql, update_values)
            skipped_count += 1

    conn.commit()
    conn.close()

    print(f'[完成] 导入完成: 新增 {imported_count} 条, 更新 {skipped_count} 条')
    return True


if __name__ == '__main__':
    if len(sys.argv) < 2:
        print('用法: python import_weekly_data.py <Excel文件路径> [店铺代码]')
        print('示例: python import_weekly_data.py "Layer2_Working/PNB-商品数据 20260518-20260524.xlsx" PNB')
        sys.exit(1)

    filepath = sys.argv[1]
    shop_code = sys.argv[2] if len(sys.argv) > 2 else None

    # 如果是相对路径，拼接
    if not os.path.isabs(filepath):
        filepath = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), filepath)

    import_excel_to_db(filepath, shop_code)
