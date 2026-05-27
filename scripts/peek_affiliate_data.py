"""
临时脚本：查看联盟数据Excel的结构
"""
import openpyxl
import os
import sys

def peek(filepath):
    wb = openpyxl.load_workbook(filepath)
    ws = wb.active
    print(f'Sheet: {ws.title}, Rows: {ws.max_row}, Cols: {ws.max_column}')
    print()
    
    # Print all rows
    for row_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True), 1):
        values = []
        for v in row:
            if v is not None:
                values.append(repr(v))
        print(f'Row {row_idx}: {values}')
    
    wb.close()

if __name__ == '__main__':
    base = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    for fn in os.listdir(os.path.join(base, 'Layer2_Working')):
        if fn.startswith('ListProducts') and fn.endswith('.xlsx'):
            print(f'\n===== {fn} =====')
            peek(os.path.join(base, 'Layer2_Working', fn))
