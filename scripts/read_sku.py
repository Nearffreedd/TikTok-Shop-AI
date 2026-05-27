import openpyxl, json, sys

wb = openpyxl.load_workbook('Layer2_Working/星昊科技有限公司SKU生成.xlsx', data_only=True)
data = {}
for sname in wb.sheetnames:
    ws = wb[sname]
    rows = []
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 80), values_only=True):
        rows.append([str(v) if v is not None else '' for v in row])
    data[sname] = rows

with open('Layer2_Working/sku_data.json', 'w', encoding='utf-8') as f:
    json.dump(data, f, ensure_ascii=False, indent=2)

print("Done. Sheets:", list(data.keys()))
for sname, rows in data.items():
    print(f"\n=== {sname} === ({len(rows)} rows)")
    for r in rows[:5]:
        print(r)
